"""Excel analyst pack built from the real CMS data: inputs in blue, every derived number a live formula.

Run:  PYTHONPATH=src python -m partd.excel_pack   ->  reports/partd_analyst_pack.xlsx
(then recalculate in LibreOffice or open in Excel; see Makefile `excel`)
"""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config
from . import public_cms as P
from . import public_extra as X

F = "Arial"
BLUE = Font(name=F, size=10, color="0000FF")
BLACK = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
HEAD = Font(name=F, size=10, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F3864")
YEL = PatternFill("solid", fgColor="FFFF00")
TITLE = Font(name=F, size=14, bold=True)
NOTE = Font(name=F, size=9, italic=True, color="52514E")
USD0 = '$#,##0;($#,##0);-'
USD2 = '$#,##0.00;($#,##0.00);-'
PCT = '0.0%;(0.0%);-'
NUM = '#,##0;(#,##0);-'


def _head(ws, row, labels, col=1):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l)
        c.font, c.fill = HEAD, HFILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 32


def _put(ws, ref, value, font=BLACK, fmt=None, fill=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build(path: str | None = None) -> str:
    d = P.build_data()
    wb = Workbook()

    # ------------------------------------------------------------------ README
    ws = wb.active
    ws.title = "README"
    _put(ws, "A1", "Medicare Part D analyst pack (real CMS public data)", TITLE)
    lines = [
        "Source: CMS data.cms.gov: Medicare Part D Spending by Drug (2020-2024), Medicare Monthly Enrollment, Part D Prescribers by Geography and Drug (2024), CMS 2026 negotiated-price fact sheet.",
        "Gross drug cost = Medicare + plan + beneficiary payments, before rebates. All of Part D, not one plan. Not UnitedHealth Group data.",
        "Blue cells are inputs copied from the CMS files by src/partd/excel_pack.py. Black cells are formulas. Yellow cells are assumptions you can change.",
        "Sheets: National_PMPM (per-member cost and a use-vs-price split), Drug_PVM (price/volume/mix by drug, 2023 to 2024), Negotiation (2026 prices and exposure),",
        "Net_Sensitivity (gross to net under an assumed rebate rate), State_PMPM (state comparison), Top_Drugs, Data_Checks.",
        "Known limits: dose units mix tablets, mL and pens; 2025 and Q1 2026 are preliminary and not used in growth rates; state cost is by prescriber location, enrollment by member residence.",
        "The rebate rate on Net_Sensitivity is an assumption, not a number from CMS data (CMS does not publish rebates).",
    ]
    for i, t in enumerate(lines, 3):
        _put(ws, f"A{i}", t, BLACK)
    ws.column_dimensions["A"].width = 160

    # ------------------------------------------------------------------ National_PMPM
    ws = wb.create_sheet("National_PMPM")
    _put(ws, "A1", "National gross drug cost per member per month", TITLE)
    _put(ws, "A2", "Inputs: CMS spending file (spend, claims) and CMS Monthly Enrollment (member months = Part D enrollees summed over 12 months).", NOTE)
    _head(ws, 4, ["Year", "Gross spend ($)", "Claims", "Member months", "Gross PMPM ($)", "Claims per 1,000 member months", "Spend per claim ($)",
                  "PMPM change vs prior year ($)", "PMPM change %", "Of which: claims per member ($)", "Of which: cost per claim ($)"])
    nat = d["natl"]
    for i, r in enumerate(nat.itertuples(), 5):
        _put(ws, f"A{i}", str(int(r.year)), BLACK)
        _put(ws, f"B{i}", float(r.spend), BLUE, USD0)
        _put(ws, f"C{i}", float(r.claims), BLUE, NUM)
        _put(ws, f"D{i}", float(r.member_months), BLUE, NUM)
        _put(ws, f"E{i}", f"=B{i}/D{i}", BLACK, USD2)
        _put(ws, f"F{i}", f"=C{i}/D{i}*1000", BLACK, '#,##0.0')
        _put(ws, f"G{i}", f"=B{i}/C{i}", BLACK, USD2)
        if i > 5:
            _put(ws, f"H{i}", f"=E{i}-E{i-1}", BLACK, USD2)
            _put(ws, f"I{i}", f"=E{i}/E{i-1}-1", BLACK, PCT)
            # exact two-factor split (log-share allocation of the PMPM change)
            _put(ws, f"J{i}", f"=H{i}*LN(F{i}/F{i-1})/(LN(F{i}/F{i-1})+LN(G{i}/G{i-1}))", BLACK, USD2)
            _put(ws, f"K{i}", f"=H{i}*LN(G{i}/G{i-1})/(LN(F{i}/F{i-1})+LN(G{i}/G{i-1}))", BLACK, USD2)
    last = 4 + len(nat)
    _put(ws, f"A{last+2}", "CAGR 2020-2024: PMPM", BOLD)
    _put(ws, f"E{last+2}", f"=(E{last}/E5)^(1/(A{last}-A5))-1", BLACK, PCT)
    _put(ws, f"A{last+3}", "CAGR: spend", BOLD)
    _put(ws, f"B{last+3}", f"=(B{last}/B5)^(1/(A{last}-A5))-1", BLACK, PCT)
    _put(ws, f"A{last+4}", "CAGR: claims", BOLD)
    _put(ws, f"C{last+4}", f"=(C{last}/C5)^(1/(A{last}-A5))-1", BLACK, PCT)
    _put(ws, f"A{last+6}", "Check: J + K equals H each year (should be 0)", NOTE)
    for i in range(6, last + 1):
        _put(ws, f"L{i}", f"=ROUND(J{i}+K{i}-H{i},6)", BLACK, '0.000000')
    _put(ws, "L4", "Split check", HEAD, fill=HFILL)
    _widths(ws, [10, 18, 16, 16, 14, 16, 14, 16, 12, 16, 16, 12])
    ws.freeze_panes = "B5"

    # ------------------------------------------------------------------ Drug_PVM (top 100 + other)
    ws = wb.create_sheet("Drug_PVM")
    _put(ws, "A1", "Price, volume and mix by drug, 2023 to 2024 (top 100 drugs by 2024 spend, all others combined)", TITLE)
    _put(ws, "A2", "Volume = dose units. Method as in the dashboard: volume = (U1/U0-1)*S0; mix = (u1 - u0*U1/U0)*p0; price = u1*(p1-p0); new drugs = u1*p1. 'All other' is treated as one drug, so mix and price inside it are not shown; totals therefore differ from the dashboard's full drug-level split (dashboard: price about -$29B).", NOTE)
    l = d["long"]
    a = l[l["year"] == 2023].set_index(["Brnd_Name", "Gnrc_Name"])
    b = l[l["year"] == 2024].set_index(["Brnd_Name", "Gnrc_Name"])
    top = b.sort_values("spend", ascending=False).head(100)
    idx = top.index
    rows = []
    for k in idx:
        rows.append((k[0], float(a["spend"].get(k, 0.0)), float(a["units"].get(k, 0.0)), float(b.loc[k, "spend"]), float(b.loc[k, "units"])))
    oth_a_s = float(a["spend"].sum() - sum(r[1] for r in rows)); oth_a_u = float(a["units"].sum() - sum(r[2] for r in rows))
    oth_b_s = float(b["spend"].sum() - sum(r[3] for r in rows)); oth_b_u = float(b["units"].sum() - sum(r[4] for r in rows))
    rows.append(("All other drugs", oth_a_s, oth_a_u, oth_b_s, oth_b_u))
    _head(ws, 4, ["Drug", "Spend 2023 ($)", "Dose units 2023", "Spend 2024 ($)", "Dose units 2024", "Price 2023 ($/unit)", "Price 2024 ($/unit)",
                  "Volume effect ($)", "Mix effect ($)", "Price effect ($)", "New-drug effect ($)", "Total change ($)", "Check vs actual"])
    first, lastr = 5, 4 + len(rows)
    for i, r in enumerate(rows, first):
        _put(ws, f"A{i}", r[0], BLACK)
        for col, v, f in (("B", r[1], USD0), ("C", r[2], NUM), ("D", r[3], USD0), ("E", r[4], NUM)):
            _put(ws, f"{col}{i}", v, BLUE, f)
        _put(ws, f"F{i}", f"=IF(C{i}>0,B{i}/C{i},0)", BLACK, '$#,##0.0000')
        _put(ws, f"G{i}", f"=IF(E{i}>0,D{i}/E{i},0)", BLACK, '$#,##0.0000')
        _put(ws, f"H{i}", f"=IF(C{i}>0,(C{i}*($E${lastr+2}/$C${lastr+2}-1)*F{i}),0)", BLACK, USD0)  # drug's share of volume effect at its own prior price
        _put(ws, f"I{i}", f"=IF(C{i}>0,(E{i}-C{i}*$E${lastr+2}/$C${lastr+2})*F{i},0)", BLACK, USD0)
        _put(ws, f"J{i}", f"=IF(AND(C{i}>0,E{i}>0),E{i}*(G{i}-F{i}),0)", BLACK, USD0)
        _put(ws, f"K{i}", f"=IF(C{i}=0,D{i},0)", BLACK, USD0)
        _put(ws, f"L{i}", f"=SUM(H{i}:K{i})", BLACK, USD0)
        _put(ws, f"M{i}", f"=ROUND(L{i}-(D{i}-B{i}),0)", BLACK, '0')
    t = lastr + 2
    _put(ws, f"A{t}", "Total", BOLD)
    for col, f in (("B", USD0), ("C", NUM), ("D", USD0), ("E", NUM), ("H", USD0), ("I", USD0), ("J", USD0), ("K", USD0), ("L", USD0), ("M", '0')):
        _put(ws, f"{col}{t}", f"=SUM({col}{first}:{col}{lastr})", BOLD, f)
    _put(ws, f"A{t+2}", "Share of total change: price", BLACK); _put(ws, f"J{t+2}", f"=J{t}/L{t}", BLACK, PCT)
    _put(ws, f"A{t+3}", "Share of total change: volume + mix", BLACK); _put(ws, f"J{t+3}", f"=(H{t}+I{t})/L{t}", BLACK, PCT)
    _put(ws, f"A{t+5}", "Note: row totals use drug-level volume/mix at each drug's own prior price; the identity (column M) should read 0 on every row.", NOTE)
    _widths(ws, [30, 18, 16, 18, 16, 14, 14, 16, 16, 16, 16, 16, 12])
    ws.freeze_panes = "B5"

    # ------------------------------------------------------------------ Negotiation
    ws = wb.create_sheet("Negotiation")
    _put(ws, "A1", "Medicare negotiated prices for 2026 (first 10 drugs): exposure and early check", TITLE)
    _put(ws, "A2", f"List and negotiated prices (30-day) transcribed from the CMS fact sheet: {X.MFP_SOURCE}", NOTE)
    _head(ws, 4, ["Drug", "List price, 30 days ($)", "Negotiated price, 30 days ($)", "Cut", "2024 gross spend ($)", "Share of 2024 Part D spend",
                  "Cut applied to 2024 spend ($)", "Spend per claim 2025 ($)", "Spend per claim Q1 2026 ($)", "Observed change"])
    n = d["negotiation"].sort_values("spend_2024", ascending=False).reset_index(drop=True)
    for i, r in enumerate(n.itertuples(), 5):
        _put(ws, f"A{i}", r.drug, BLACK)
        _put(ws, f"B{i}", float(r.list_price_30d), BLUE, USD2)
        _put(ws, f"C{i}", float(r.mfp_30d), BLUE, USD2)
        _put(ws, f"D{i}", f"=1-C{i}/B{i}", BLACK, PCT)
        _put(ws, f"E{i}", float(r.spend_2024), BLUE, USD0)
        _put(ws, f"F{i}", f"=E{i}/National_PMPM!$B$9", BLACK, PCT)
        _put(ws, f"G{i}", f"=E{i}*D{i}", BLACK, USD0)
        _put(ws, f"H{i}", float(r.spend_per_claim_2025), BLUE, USD2)
        _put(ws, f"I{i}", float(r.spend_per_claim_q1_2026), BLUE, USD2)
        _put(ws, f"J{i}", f"=I{i}/H{i}-1", BLACK, PCT)
    e = 4 + len(n)
    _put(ws, f"A{e+1}", "Total", BOLD)
    for col, f in (("E", USD0), ("F", PCT), ("G", USD0)):
        _put(ws, f"{col}{e+1}", f"=SUM({col}5:{col}{e})", BOLD, f)
    _put(ws, f"A{e+3}", "Cut as % of these drugs' spend", BLACK); _put(ws, f"G{e+3}", f"=G{e+1}/E{e+1}", BLACK, PCT)
    _put(ws, f"A{e+4}", "Per member per month at 2024 volume ($)", BLACK); _put(ws, f"G{e+4}", f"=G{e+1}/National_PMPM!$D$9", BLACK, USD2)
    _put(ws, f"A{e+6}", "Constant-volume, gross, vs list: a ceiling, not a forecast. Rebates already paid on these drugs make the plan's net effect smaller. NovoLog's list price was already cut in 2024, so its observed drop is small.", NOTE)
    _widths(ws, [18, 16, 18, 10, 18, 16, 20, 16, 16, 12])

    # ------------------------------------------------------------------ Net_Sensitivity
    ws = wb.create_sheet("Net_Sensitivity")
    _put(ws, "A1", "Gross to net: what an assumed rebate rate does to cost per member per month", TITLE)
    _put(ws, "A2", "CMS does not publish rebates. The rate below is an ASSUMPTION for illustration, not a number from the data. Change the yellow cell.", NOTE)
    _put(ws, "A4", "Assumed rebate rate (share of gross spend)", BOLD)
    _put(ws, "B4", 0.20, BLUE, PCT, YEL)
    _put(ws, "A5", "2024 gross PMPM ($)", BLACK); _put(ws, "B5", "=National_PMPM!E9", BLACK, USD2)
    _put(ws, "A6", "Implied net PMPM ($)", BLACK); _put(ws, "B6", "=B5*(1-B4)", BLACK, USD2)
    _put(ws, "A7", "Implied net 2024 spend ($)", BLACK); _put(ws, "B7", "=National_PMPM!B9*(1-B4)", BLACK, USD0)
    _head(ws, 9, ["Rebate rate", "Net PMPM ($)", "Net 2024 spend ($)"])
    for j, r in enumerate([0, .05, .10, .15, .20, .25, .30, .35, .40], 10):
        _put(ws, f"A{j}", r, BLUE, PCT)
        _put(ws, f"B{j}", f"=$B$5*(1-A{j})", BLACK, USD2)
        _put(ws, f"C{j}", f"=National_PMPM!$B$9*(1-A{j})", BLACK, USD0)
    _widths(ws, [44, 18, 20])

    # ------------------------------------------------------------------ State_PMPM
    ws = wb.create_sheet("State_PMPM")
    _put(ws, "A1", "Gross drug cost per member per month by state, 2024", TITLE)
    _put(ws, "A2", "Cost is by prescriber location; member months by residence. DC is inflated by cross-border prescribing. Territories excluded.", NOTE)
    _head(ws, 4, ["State", "Gross spend ($)", "Member months", "30-day fills", "Gross PMPM ($)", "vs national", "Rank (1 = highest)", "Fills per member per month", "Cost per fill ($)"])
    s = d["states"].sort_values("state").reset_index(drop=True)
    for i, r in enumerate(s.itertuples(), 5):
        _put(ws, f"A{i}", r.state, BLACK)
        _put(ws, f"B{i}", float(r.spend), BLUE, USD0)
        _put(ws, f"C{i}", float(r.member_months), BLUE, NUM)
        _put(ws, f"D{i}", float(r.fills30), BLUE, '#,##0.0')
        _put(ws, f"E{i}", f"=B{i}/C{i}", BLACK, USD2)
        _put(ws, f"F{i}", f"=E{i}/National_PMPM!$E$9-1", BLACK, PCT)
        _put(ws, f"G{i}", f"=RANK(E{i},$E$5:$E${4+len(s)})", BLACK, '0')
        _put(ws, f"H{i}", f"=D{i}/C{i}", BLACK, '0.00')
        _put(ws, f"I{i}", f"=B{i}/D{i}", BLACK, USD2)
    ws.conditional_formatting.add(f"E5:E{4+len(s)}", ColorScaleRule(start_type="min", start_color="DCE9F9", end_type="max", end_color="2A78D6"))
    _widths(ws, [10, 18, 16, 14, 14, 12, 14, 18, 14])
    ws.freeze_panes = "B5"

    # ------------------------------------------------------------------ Top_Drugs
    ws = wb.create_sheet("Top_Drugs")
    _put(ws, "A1", "Top 25 drugs by 2024 gross spend", TITLE)
    _head(ws, 3, ["Rank", "Drug", "Spend 2020 ($)", "Spend 2024 ($)", "Share of 2024", "Cumulative share", "Change 2020-2024 ($)", "Claims 2024", "Spend per claim 2024 ($)"])
    tp = b.sort_values("spend", ascending=False).head(25)
    a20 = l[l["year"] == 2020].set_index(["Brnd_Name", "Gnrc_Name"])
    for i, (k, r) in enumerate(tp.iterrows(), 4):
        _put(ws, f"A{i}", i - 3, BLACK, '0')
        _put(ws, f"B{i}", k[0], BLACK)
        _put(ws, f"C{i}", float(a20["spend"].get(k, 0.0)), BLUE, USD0)
        _put(ws, f"D{i}", float(r["spend"]), BLUE, USD0)
        _put(ws, f"E{i}", f"=D{i}/National_PMPM!$B$9", BLACK, PCT)
        _put(ws, f"F{i}", f"=SUM($E$4:E{i})", BLACK, PCT)
        _put(ws, f"G{i}", f"=D{i}-C{i}", BLACK, USD0)
        _put(ws, f"H{i}", float(r["claims"]), BLUE, NUM)
        _put(ws, f"I{i}", f"=D{i}/H{i}", BLACK, USD2)
    _widths(ws, [8, 26, 18, 18, 14, 16, 20, 16, 18])

    # ------------------------------------------------------------------ Data_Checks
    ws = wb.create_sheet("Data_Checks")
    _put(ws, "A1", "Cross-file reconciliation", TITLE)
    _head(ws, 3, ["Check", "A", "B", "Difference", "Tolerance", "Result"])
    for i, c in enumerate(d["recon"], 4):
        _put(ws, f"A{i}", c["check"], BLACK)
        _put(ws, f"B{i}", float(c["a"]), BLUE, NUM)
        _put(ws, f"C{i}", float(c["b"]), BLUE, NUM)
        if c["tolerance"] is None:
            _put(ws, f"D{i}", f"=B{i}/C{i}", BLACK, PCT)
            _put(ws, f"E{i}", "info", BLACK)
            _put(ws, f"F{i}", "info", BLACK)
        elif c["tolerance"] == 0:
            _put(ws, f"D{i}", f"=B{i}-C{i}", BLACK, NUM)
            _put(ws, f"E{i}", 0, BLUE, NUM)
            _put(ws, f"F{i}", f'=IF(D{i}=E{i},"pass","FAIL")', BLACK)
        else:
            _put(ws, f"D{i}", f"=IF(B{i}=0,0,C{i}/B{i}-1)", BLACK, '0.000%')
            _put(ws, f"E{i}", float(c["tolerance"]), BLUE, '0.000%')
            _put(ws, f"F{i}", f'=IF(ABS(D{i})<=E{i},"pass","FAIL")', BLACK)
    _widths(ws, [66, 20, 20, 14, 12, 10])

    out = path or str(config.REPORT_DIR / "partd_analyst_pack.xlsx")
    config.REPORT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print(build())
