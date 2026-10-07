"""Tests for the per-member, state and negotiation work on real CMS data (skipped until `make public` has run)."""
import numpy as np
import openpyxl
import pytest

from partd import public_cms as P
from partd import public_extra as X

pytestmark = pytest.mark.skipif(not (P.ANNUAL.exists() and X.ENROLL.exists() and X.GEO.exists()), reason="run `make public` first")


@pytest.fixture(scope="module")
def d():
    return P.build_data()


def test_member_months_complete_years(d):
    mm = d["member_months"].set_index("year")
    assert (mm.loc[2020:2024, "months"] == 12).all()
    assert 45e6 < mm.loc[2020, "avg_enrollees"] < 60e6


def test_national_gross_pmpm_plausible_and_consistent(d):
    n = d["natl"].set_index("year")
    assert 300 < n.loc[2020, "gross_pmpm"] < 400 and 400 < n.loc[2024, "gross_pmpm"] < 500
    assert np.allclose(n["gross_pmpm"], n["claims_per_1000_mm"] / 1000 * n["spend_per_claim"])


def test_pmpm_bridge_sums_exactly(d):
    for r in d["bridge"].itertuples():
        assert abs(r.utilization + r.cost_per_claim - r.change) < 1e-9


def test_state_table_covers_states_and_dc(d):
    s = d["states"]
    assert len(s) == 51 and s["gross_pmpm"].between(150, 1000).all()
    # 50 states + DC cover ~98.7% of national spend; the rest is Puerto Rico and other territories, excluded on purpose
    assert 0.97 < s["spend"].sum() / d["totals"].set_index("year").loc[2024, "spend"] < 1.0


def test_reconciliation_checks_pass(d):
    assert all(c["pass"] for c in d["recon"])


def test_negotiation_table_matches_cms_sheet(d):
    n = d["negotiation"].set_index("drug")
    assert len(n) == 10
    assert abs(n.loc["Eliquis", "mfp_discount"] - (1 - 231 / 521)) < 1e-9
    assert 0.15 < n["spend_2024"].sum() / d["totals"].set_index("year").loc[2024, "spend"] < 0.30
    assert (n["observed_change_per_claim"] < 0).all()


def test_excel_pack_is_formula_driven(tmp_path, d):
    from partd import excel_pack
    path = excel_pack.build(str(tmp_path / "pack.xlsx"))
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[:2] == ["README", "National_PMPM"]
    nat = wb["National_PMPM"]
    assert nat["E9"].value == "=B9/D9"  # PMPM is a formula, not a pasted number
    formulas = sum(1 for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    assert formulas > 800


def test_briefing_builds(d):
    from partd import briefing
    text = briefing.build()
    assert "## Summary" in text and "Limitations" in text
